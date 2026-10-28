from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.shortcuts import get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.utils import timezone

from accounts.mixins import BranchPermissionMixin, ProtectedDeleteMixin, filter_by_branch
from courses.models import Course
from students.models import Student
from prospects.models import Prospect
from prospects.utils import get_user_root_branch_ids, is_root_branch, do_convert_prospect_to_student
from .models import Account, AttachType, Attach, AccountAttach, AccountCondition, AccountNote
from .forms import (
    AccountForm, AttachTypeForm, AttachForm,
    AccountAttachForm, AccountConditionForm, AccountNoteForm,
)


def _apply_account_filters(queryset, params):
    """Apply shared filter logic for Account (Registration) list and export."""
    search = params.get('search')
    if search:
        queryset = queryset.filter(
            Q(student__contact__first_name__icontains=search) |
            Q(student__contact__forth_name__icontains=search)
        )
    course = params.get('course')
    if course:
        queryset = queryset.filter(course_id=course)
    student = params.get('student')
    if student:
        queryset = queryset.filter(student_id=student)
    payment_type = params.get('payment_type')
    if payment_type:
        queryset = queryset.filter(course_payment_type=payment_type)
    return queryset


# ============================================================
# Account Views
# ============================================================

class AccountListView(BranchPermissionMixin, ListView):
    model = Account
    template_name = 'registrations/account_list.html'
    context_object_name = 'accounts'
    paginate_by = 20
    required_perm = 'view_registration'
    branch_field = 'course__master__branch'

    def get_queryset(self):
        queryset = Account.objects.select_related('student', 'course', 'course__master', 'last_person').all()
        queryset = filter_by_branch(queryset, self.request.user, self.branch_field, perm=self.required_perm)
        return _apply_account_filters(queryset, self.request.GET)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['courses'] = Course.objects.select_related('master').all()
        context['students'] = Student.objects.select_related('contact').all()
        root_ids = get_user_root_branch_ids(self.request.user, 'add_registration')
        context['prospects'] = Prospect.objects.filter(
            student__isnull=True, branch_id__in=root_ids
        ).order_by('-created_at')
        context['is_root_company'] = bool(root_ids)
        context['pending_registered_prospects'] = Prospect.objects.filter(
            branch_id__in=root_ids, status='registered', student__isnull=True
        ).select_related('master', 'branch').order_by('-updated_at')
        return context


class AccountDetailView(BranchPermissionMixin, DetailView):
    model = Account
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    template_name = 'registrations/account_detail.html'
    context_object_name = 'account'
    required_perm = 'view_registration'
    branch_field = 'course__master__branch'

    def get_queryset(self):
        return Account.objects.select_related('student', 'course', 'course__master', 'last_person')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['courses'] = Course.objects.select_related('master').all()
        context['students'] = Student.objects.select_related('contact').all()
        return context


class AccountCreateView(BranchPermissionMixin, CreateView):
    model = Account
    form_class = AccountForm
    template_name = 'registrations/account_form.html'
    success_url = reverse_lazy('registration-list')
    required_perm = 'add_registration'
    branch_field = 'course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.last_person = self.request.user
        prospect = form.cleaned_data.get('prospect')
        course = form.cleaned_data.get('course')
        if prospect and course and course.master and course.master.branch and is_root_branch(course.master.branch):
            student = do_convert_prospect_to_student(prospect, self.request.user)
            form.instance.student = student
        return super().form_valid(form)


class AccountUpdateView(BranchPermissionMixin, UpdateView):
    model = Account
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    form_class = AccountForm
    template_name = 'registrations/account_form.html'
    success_url = reverse_lazy('registration-list')
    required_perm = 'change_registration'
    branch_field = 'course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.last_person = self.request.user
        return super().form_valid(form)


class AccountDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = Account
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    template_name = 'registrations/account_confirm_delete.html'
    success_url = reverse_lazy('registration-list')
    required_perm = 'delete_registration'
    branch_field = 'course__master__branch'


@login_required
@require_POST
def account_create_ajax(request):
    if not request.user.has_perm_on_any_branch('add_registration'):
        return JsonResponse({'success': False, 'message': 'غير مسموح لك دخول هنا'}, status=403)
    form = AccountForm(request.POST, user=request.user)
    if form.is_valid():
        account = form.save(commit=False)
        account.last_person = request.user

        # في شركة الجذور، المستفسر بيتحول لطالب قبل التسجيل
        prospect = form.cleaned_data.get('prospect')
        course = form.cleaned_data.get('course')
        if prospect and course and course.master and course.master.branch and is_root_branch(course.master.branch):
            student = do_convert_prospect_to_student(prospect, request.user)
            account.student = student

        account.save()
        return JsonResponse({'success': True, 'message': 'تم إنشاء التسجيل بنجاح', 'id': account.id, 'slug': account.slug})
    return JsonResponse({'success': False, 'errors': form.errors}, status=400)


@login_required
@require_POST
def account_update_ajax(request, pk):
    account = get_object_or_404(Account, pk=pk)
    if not request.user.has_perm('change_registration', branch=account.course.master.branch):
        return JsonResponse({'success': False, 'message': 'غير مسموح لك دخول هنا'}, status=403)
    form = AccountForm(request.POST, instance=account, user=request.user)
    if form.is_valid():
        account = form.save(commit=False)
        account.last_person = request.user
        account.save()
        return JsonResponse({'success': True, 'message': 'تم تحديث التسجيل بنجاح'})
    return JsonResponse({'success': False, 'errors': form.errors}, status=400)


# ============================================================
# Bulk Export (filtered list)
# ============================================================

@login_required
def export_registrations_excel(request):
    """Export filtered registrations list as Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.cell.cell import MergedCell

    user = request.user
    qs = filter_by_branch(
        Account.objects.select_related('student', 'course', 'course__master'),
        user, 'course__master__branch', perm='view_registration'
    )
    qs = _apply_account_filters(qs, request.GET)
    qs = qs.order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'التسجيلات'
    ws.sheet_view.rightToLeft = True

    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    headers = ['المفتاح', 'الطالب', 'الدورة', 'نوع الدفع', 'السعر', 'المطلوب', 'المدفوع', 'المتبقي']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[1].height = 25

    for row_num, account in enumerate(qs, 2):
        values = [
            account.get_key(),
            account.student.get_full_name(),
            f'{account.course.master.name} - {account.course.code}',
            account.course_payment_type,
            float(account.course_price),
            float(account.get_required_price()),
            float(account.get_paid_price()),
            float(account.get_remain_price()),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    for col_cells in ws.columns:
        if isinstance(col_cells[0], MergedCell):
            continue
        max_len = 0
        for cell in col_cells:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

    now = timezone.localtime().strftime('%Y%m-%d_%H%M')
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="registrations_report_{now}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_registrations_pdf(request):
    """Export filtered registrations list as a PDF report table."""
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_RIGHT
    from offers.views import ARABIC_FONT, _prepare_arabic

    user = request.user
    qs = filter_by_branch(
        Account.objects.select_related('student', 'course', 'course__master'),
        user, 'course__master__branch', perm='view_registration'
    )
    qs = _apply_account_filters(qs, request.GET)
    qs = qs.order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm,
                            topMargin=1*cm, bottomMargin=1*cm)
    elements = []

    font_name = ARABIC_FONT
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleAr', parent=styles['Normal'], fontName=font_name, fontSize=16, alignment=TA_RIGHT, spaceAfter=12, leading=22)
    header_style = ParagraphStyle('HeaderAr', parent=styles['Normal'], fontName=font_name, fontSize=9, alignment=TA_RIGHT,
                                  textColor=colors.white, leading=14)
    cell_style = ParagraphStyle('CellAr', parent=styles['Normal'], fontName=font_name, fontSize=8, alignment=TA_RIGHT, leading=12)

    elements.append(Paragraph(_prepare_arabic('تقرير التسجيلات'), title_style))
    elements.append(Spacer(1, 0.3*cm))

    headers = ['المفتاح', 'الطالب', 'الدورة', 'نوع الدفع', 'السعر', 'المطلوب', 'المدفوع', 'المتبقي']
    table_data = [[Paragraph(_prepare_arabic(h), header_style) for h in headers]]

    for account in qs:
        row = [
            Paragraph(_prepare_arabic(account.get_key()), cell_style),
            Paragraph(_prepare_arabic(account.student.get_full_name()), cell_style),
            Paragraph(_prepare_arabic(f'{account.course.master.name} - {account.course.code}'), cell_style),
            Paragraph(_prepare_arabic(account.course_payment_type), cell_style),
            Paragraph(_prepare_arabic(f'{account.course_price}'), cell_style),
            Paragraph(_prepare_arabic(f'{account.get_required_price()}'), cell_style),
            Paragraph(_prepare_arabic(f'{account.get_paid_price()}'), cell_style),
            Paragraph(_prepare_arabic(f'{account.get_remain_price()}'), cell_style),
        ]
        table_data.append(row)

    page_w = landscape(A4)[0] - 2*cm
    col_widths = [page_w*0.13, page_w*0.20, page_w*0.20, page_w*0.11, page_w*0.09, page_w*0.09, page_w*0.09, page_w*0.09]

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor('#2563EB')),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    count_style = ParagraphStyle('CountAr', parent=styles['Normal'], fontName=font_name, fontSize=10, alignment=TA_RIGHT)
    elements.append(Paragraph(_prepare_arabic(f'إجمالي النتائج: {qs.count()}'), count_style))

    doc.build(elements)

    now = timezone.localtime().strftime('%Y%m-%d_%H%M')
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="registrations_report_{now}.pdf"'
    return response


# ============================================================
# AttachType Views
# ============================================================

class AttachTypeListView(BranchPermissionMixin, ListView):
    model = AttachType
    template_name = 'registrations/attachtype_list.html'
    context_object_name = 'attach_types'
    paginate_by = 20
    required_perm = 'view_attachtype'


class AttachTypeDetailView(BranchPermissionMixin, DetailView):
    model = AttachType
    template_name = 'registrations/attachtype_detail.html'
    context_object_name = 'attach_type'
    required_perm = 'view_attachtype'


class AttachTypeCreateView(BranchPermissionMixin, CreateView):
    model = AttachType
    form_class = AttachTypeForm
    template_name = 'registrations/attachtype_form.html'
    success_url = reverse_lazy('attachtype-list')
    required_perm = 'add_attachtype'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class AttachTypeUpdateView(BranchPermissionMixin, UpdateView):
    model = AttachType
    form_class = AttachTypeForm
    template_name = 'registrations/attachtype_form.html'
    success_url = reverse_lazy('attachtype-list')
    required_perm = 'change_attachtype'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class AttachTypeDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = AttachType
    template_name = 'registrations/attachtype_confirm_delete.html'
    success_url = reverse_lazy('attachtype-list')
    required_perm = 'delete_attachtype'


# ============================================================
# Attach Views
# ============================================================

class AttachListView(BranchPermissionMixin, ListView):
    model = Attach
    template_name = 'registrations/attach_list.html'
    context_object_name = 'attaches'
    paginate_by = 20
    required_perm = 'view_attach'
    branch_field = 'person__branch'

    def get_queryset(self):
        queryset = Attach.objects.select_related('attach_type', 'person').all()
        queryset = filter_by_branch(queryset, self.request.user, self.branch_field, perm=self.required_perm)
        attach_type = self.request.GET.get('type')
        if attach_type:
            queryset = queryset.filter(attach_type_id=attach_type)
        return queryset


class AttachDetailView(BranchPermissionMixin, DetailView):
    model = Attach
    template_name = 'registrations/attach_detail.html'
    context_object_name = 'attach'
    required_perm = 'view_attach'
    branch_field = 'person__branch'

    def get_queryset(self):
        return Attach.objects.select_related('attach_type', 'person')


class AttachCreateView(BranchPermissionMixin, CreateView):
    model = Attach
    form_class = AttachForm
    template_name = 'registrations/attach_form.html'
    success_url = reverse_lazy('attach-list')
    required_perm = 'add_attach'
    branch_field = 'person__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.person = self.request.user
        return super().form_valid(form)


class AttachUpdateView(BranchPermissionMixin, UpdateView):
    model = Attach
    form_class = AttachForm
    template_name = 'registrations/attach_form.html'
    success_url = reverse_lazy('attach-list')
    required_perm = 'change_attach'
    branch_field = 'person__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class AttachDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = Attach
    template_name = 'registrations/attach_confirm_delete.html'
    success_url = reverse_lazy('attach-list')
    required_perm = 'delete_attach'
    branch_field = 'person__branch'


# ============================================================
# AccountAttach Views
# ============================================================

class AccountAttachListView(BranchPermissionMixin, ListView):
    model = AccountAttach
    template_name = 'registrations/accountattach_list.html'
    context_object_name = 'account_attaches'
    paginate_by = 20
    required_perm = 'view_accountattach'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        queryset = AccountAttach.objects.select_related('account', 'attach').all()
        queryset = filter_by_branch(queryset, self.request.user, self.branch_field, perm=self.required_perm)
        account = self.request.GET.get('account')
        if account:
            queryset = queryset.filter(account_id=account)
        return queryset


class AccountAttachDetailView(BranchPermissionMixin, DetailView):
    model = AccountAttach
    template_name = 'registrations/accountattach_detail.html'
    context_object_name = 'account_attach'
    required_perm = 'view_accountattach'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        return AccountAttach.objects.select_related('account', 'attach')


class AccountAttachCreateView(BranchPermissionMixin, CreateView):
    model = AccountAttach
    form_class = AccountAttachForm
    template_name = 'registrations/accountattach_form.html'
    success_url = reverse_lazy('accountattach-list')
    required_perm = 'add_accountattach'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class AccountAttachUpdateView(BranchPermissionMixin, UpdateView):
    model = AccountAttach
    form_class = AccountAttachForm
    template_name = 'registrations/accountattach_form.html'
    success_url = reverse_lazy('accountattach-list')
    required_perm = 'change_accountattach'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs


class AccountAttachDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = AccountAttach
    template_name = 'registrations/accountattach_confirm_delete.html'
    success_url = reverse_lazy('accountattach-list')
    required_perm = 'delete_accountattach'
    branch_field = 'account__course__master__branch'


# ============================================================
# AccountCondition Views
# ============================================================

class AccountConditionListView(BranchPermissionMixin, ListView):
    model = AccountCondition
    template_name = 'registrations/accountcondition_list.html'
    context_object_name = 'account_conditions'
    paginate_by = 20
    required_perm = 'view_accountcondition'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        queryset = AccountCondition.objects.select_related('account', 'person').all()
        queryset = filter_by_branch(queryset, self.request.user, self.branch_field, perm=self.required_perm)
        account = self.request.GET.get('account')
        if account:
            queryset = queryset.filter(account_id=account)
        fulfilled = self.request.GET.get('fulfilled')
        if fulfilled is not None:
            queryset = queryset.filter(fulfilled=fulfilled.lower() == 'true')
        return queryset


class AccountConditionDetailView(BranchPermissionMixin, DetailView):
    model = AccountCondition
    template_name = 'registrations/accountcondition_detail.html'
    context_object_name = 'account_condition'
    required_perm = 'view_accountcondition'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        return AccountCondition.objects.select_related('account', 'person')


class AccountConditionCreateView(BranchPermissionMixin, CreateView):
    model = AccountCondition
    form_class = AccountConditionForm
    template_name = 'registrations/accountcondition_form.html'
    success_url = reverse_lazy('accountcondition-list')
    required_perm = 'add_accountcondition'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.person = self.request.user
        return super().form_valid(form)


class AccountConditionUpdateView(BranchPermissionMixin, UpdateView):
    model = AccountCondition
    form_class = AccountConditionForm
    template_name = 'registrations/accountcondition_form.html'
    success_url = reverse_lazy('accountcondition-list')
    required_perm = 'change_accountcondition'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.person = self.request.user
        return super().form_valid(form)


class AccountConditionDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = AccountCondition
    template_name = 'registrations/accountcondition_confirm_delete.html'
    success_url = reverse_lazy('accountcondition-list')
    required_perm = 'delete_accountcondition'
    branch_field = 'account__course__master__branch'


# ============================================================
# AccountNote Views
# ============================================================

class AccountNoteListView(BranchPermissionMixin, ListView):
    model = AccountNote
    template_name = 'registrations/accountnote_list.html'
    context_object_name = 'account_notes'
    paginate_by = 20
    required_perm = 'view_accountnote'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        queryset = AccountNote.objects.select_related('account', 'person').all()
        queryset = filter_by_branch(queryset, self.request.user, self.branch_field, perm=self.required_perm)
        account = self.request.GET.get('account')
        if account:
            queryset = queryset.filter(account_id=account)
        return queryset


class AccountNoteDetailView(BranchPermissionMixin, DetailView):
    model = AccountNote
    template_name = 'registrations/accountnote_detail.html'
    context_object_name = 'account_note'
    required_perm = 'view_accountnote'
    branch_field = 'account__course__master__branch'

    def get_queryset(self):
        return AccountNote.objects.select_related('account', 'person')


class AccountNoteCreateView(BranchPermissionMixin, CreateView):
    model = AccountNote
    form_class = AccountNoteForm
    template_name = 'registrations/accountnote_form.html'
    success_url = reverse_lazy('accountnote-list')
    required_perm = 'add_accountnote'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.person = self.request.user
        return super().form_valid(form)


class AccountNoteUpdateView(BranchPermissionMixin, UpdateView):
    model = AccountNote
    form_class = AccountNoteForm
    template_name = 'registrations/accountnote_form.html'
    success_url = reverse_lazy('accountnote-list')
    required_perm = 'change_accountnote'
    branch_field = 'account__course__master__branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        form.instance.person = self.request.user
        return super().form_valid(form)


class AccountNoteDeleteView(ProtectedDeleteMixin, BranchPermissionMixin, DeleteView):
    model = AccountNote
    template_name = 'registrations/accountnote_confirm_delete.html'
    success_url = reverse_lazy('accountnote-list')
    required_perm = 'delete_accountnote'
    branch_field = 'account__course__master__branch'

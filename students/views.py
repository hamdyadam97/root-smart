from django.db.models import Q
from django.urls import reverse_lazy
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.http import JsonResponse, HttpResponse
from django.views.decorators.http import require_POST
from django.shortcuts import get_object_or_404
from django.contrib.auth.decorators import login_required
from django.utils import timezone

from accounts.mixins import BranchPermissionMixin, filter_by_branch

from .models import Contact, Student
from .forms import StudentForm


def _apply_student_filters(queryset, params):
    """Apply shared filter logic for Student list and export."""
    search = params.get('search')
    if search:
        queryset = queryset.filter(
            Q(contact__first_name__icontains=search) |
            Q(contact__forth_name__icontains=search) |
            Q(contact__mobile__icontains=search) |
            Q(slug__icontains=search)
        )
    level = params.get('level')
    if level:
        queryset = queryset.filter(level=level)
    preferred = params.get('preferred_contact')
    if preferred:
        queryset = queryset.filter(preferred_contact=preferred)
    return queryset.distinct()


class StudentListView(BranchPermissionMixin, ListView):
    model = Student
    template_name = 'students/student_list.html'
    context_object_name = 'students'
    paginate_by = 20
    required_perm = 'view_student'
    branch_field = 'branch'

    def get_queryset(self):
        queryset = Student.objects.select_related('contact', 'branch').all()
        queryset = filter_by_branch(queryset, self.request.user, 'branch', perm=self.required_perm)
        return _apply_student_filters(queryset, self.request.GET)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from core.models import Branch
        allowed_ids = [b.pk for b in self.request.user.get_branches_for_perm(self.required_perm)]
        if self.request.user.is_executive():
            context['branches'] = Branch.objects.all().order_by('code', 'name')
        else:
            context['branches'] = Branch.objects.filter(pk__in=allowed_ids).order_by('code', 'name')
        return context


class StudentDetailView(BranchPermissionMixin, DetailView):
    required_perm = 'view_student'
    model = Student
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    template_name = 'students/student_detail.html'
    context_object_name = 'student'
    branch_field = 'branch'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        from core.models import Branch
        allowed_ids = [b.pk for b in self.request.user.get_branches_for_perm(self.required_perm)]
        if self.request.user.is_executive():
            context['branches'] = Branch.objects.all().order_by('code', 'name')
        else:
            context['branches'] = Branch.objects.filter(pk__in=allowed_ids).order_by('code', 'name')
        return context


class StudentCreateView(BranchPermissionMixin, CreateView):
    required_perm = 'add_student'
    model = Student
    form_class = StudentForm
    template_name = 'students/student_form.html'
    success_url = reverse_lazy('student-list')
    branch_field = 'branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        contact_data = {
            'first_name': form.cleaned_data['first_name'],
            'second_name': form.cleaned_data['second_name'],
            'third_name': form.cleaned_data['third_name'],
            'forth_name': form.cleaned_data['forth_name'],
            'address': form.cleaned_data['address'],
            'mobile': form.cleaned_data['mobile'],
            'phone': form.cleaned_data['phone'],
            'nationality': form.cleaned_data['nationality'],
            'identity_number': form.cleaned_data['identity_number'],
            'identity_location': form.cleaned_data['identity_location'],
            'identity_start_date': form.cleaned_data['identity_start_date'],
            'birth_date': form.cleaned_data['birth_date'],
            'birth_location': form.cleaned_data['birth_location'],
            'qualification': form.cleaned_data['qualification'],
            'photo': form.cleaned_data['photo'],
        }
        contact = Contact.objects.create(**contact_data)
        form.instance.contact = contact
        return super().form_valid(form)


class StudentUpdateView(BranchPermissionMixin, UpdateView):
    required_perm = 'change_student'
    model = Student
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    form_class = StudentForm
    template_name = 'students/student_form.html'
    success_url = reverse_lazy('student-list')
    branch_field = 'branch'

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs['user'] = self.request.user
        return kwargs

    def form_valid(self, form):
        contact = self.object.contact
        contact.first_name = form.cleaned_data['first_name']
        contact.second_name = form.cleaned_data['second_name']
        contact.third_name = form.cleaned_data['third_name']
        contact.forth_name = form.cleaned_data['forth_name']
        contact.address = form.cleaned_data['address']
        contact.mobile = form.cleaned_data['mobile']
        contact.phone = form.cleaned_data['phone']
        contact.nationality = form.cleaned_data['nationality']
        contact.identity_number = form.cleaned_data['identity_number']
        contact.identity_location = form.cleaned_data['identity_location']
        contact.identity_start_date = form.cleaned_data['identity_start_date']
        contact.birth_date = form.cleaned_data['birth_date']
        contact.birth_location = form.cleaned_data['birth_location']
        contact.qualification = form.cleaned_data['qualification']
        contact.photo = form.cleaned_data['photo']
        contact.save()
        return super().form_valid(form)


class StudentDeleteView(BranchPermissionMixin, DeleteView):
    required_perm = 'delete_student'
    model = Student
    slug_field = 'slug'
    slug_url_kwarg = 'slug'
    template_name = 'students/student_confirm_delete.html'
    success_url = reverse_lazy('student-list')
    branch_field = 'branch'


@require_POST
def student_create_ajax(request):
    if not request.user.has_perm_on_any_branch('add_student'):
        return JsonResponse({'success': False, 'message': 'غير مسموح لك دخول هنا'}, status=403)
    form = StudentForm(request.POST, request.FILES, user=request.user)
    if form.is_valid():
        contact_data = {
            'first_name': form.cleaned_data['first_name'],
            'second_name': form.cleaned_data['second_name'],
            'third_name': form.cleaned_data['third_name'],
            'forth_name': form.cleaned_data['forth_name'],
            'address': form.cleaned_data['address'],
            'mobile': form.cleaned_data['mobile'],
            'phone': form.cleaned_data['phone'],
            'nationality': form.cleaned_data['nationality'],
            'identity_number': form.cleaned_data['identity_number'],
            'identity_location': form.cleaned_data['identity_location'],
            'identity_start_date': form.cleaned_data['identity_start_date'],
            'birth_date': form.cleaned_data['birth_date'],
            'birth_location': form.cleaned_data['birth_location'],
            'qualification': form.cleaned_data['qualification'],
            'photo': form.cleaned_data['photo'],
        }
        contact = Contact.objects.create(**contact_data)
        student = Student.objects.create(
            contact=contact,
            branch=form.cleaned_data['branch'],
            level=form.cleaned_data['level'],
            preferred_contact=form.cleaned_data['preferred_contact']
        )
        return JsonResponse({'success': True, 'message': 'تم إنشاء الطالب بنجاح', 'id': student.id, 'slug': student.slug})
    return JsonResponse({'success': False, 'errors': form.errors}, status=400)


@require_POST
def student_update_ajax(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if not request.user.has_perm('change_student', branch=student.branch):
        return JsonResponse({'success': False, 'message': 'غير مسموح لك دخول هنا'}, status=403)
    form = StudentForm(request.POST, request.FILES, instance=student, user=request.user)
    if form.is_valid():
        contact = student.contact
        contact.first_name = form.cleaned_data['first_name']
        contact.second_name = form.cleaned_data['second_name']
        contact.third_name = form.cleaned_data['third_name']
        contact.forth_name = form.cleaned_data['forth_name']
        contact.address = form.cleaned_data['address']
        contact.mobile = form.cleaned_data['mobile']
        contact.phone = form.cleaned_data['phone']
        contact.nationality = form.cleaned_data['nationality']
        contact.identity_number = form.cleaned_data['identity_number']
        contact.identity_location = form.cleaned_data['identity_location']
        contact.identity_start_date = form.cleaned_data['identity_start_date']
        contact.birth_date = form.cleaned_data['birth_date']
        contact.birth_location = form.cleaned_data['birth_location']
        contact.qualification = form.cleaned_data['qualification']
        contact.photo = form.cleaned_data['photo']
        contact.save()
        student.branch = form.cleaned_data['branch']
        student.level = form.cleaned_data['level']
        student.preferred_contact = form.cleaned_data['preferred_contact']
        student.save()
        return JsonResponse({'success': True, 'message': 'تم تحديث الطالب بنجاح'})
    return JsonResponse({'success': False, 'errors': form.errors}, status=400)


# ============================================================
# Bulk Export (filtered list)
# ============================================================

@login_required
def export_students_excel(request):
    """Export filtered students list as Excel."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.cell.cell import MergedCell

    user = request.user
    qs = filter_by_branch(
        Student.objects.select_related('contact', 'branch'), user, 'branch', perm='view_student'
    )
    qs = _apply_student_filters(qs, request.GET)
    qs = qs.order_by('-created_at')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'الطلاب'

    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    headers = ['الاسم', 'الجوال', 'المستوى', 'طريقة التواصل', 'الفرع', 'تاريخ الإنشاء']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws.row_dimensions[1].height = 25

    for row_num, student in enumerate(qs, 2):
        values = [
            student.get_full_name(),
            student.get_mobile() or '-',
            student.level,
            student.get_preferred_contact_display(),
            str(student.branch),
            student.created_at.strftime('%Y-%m-%d %H:%M') if student.created_at else '',
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_num, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    for col_cells in ws.columns:
        if isinstance(col_cells[0], MergedCell):
            continue
        max_len = 0
        for cell in col_cells:
            try:
                if len(str(cell.value)) > max_len:
                    max_len = len(str(cell.value))
            except Exception:
                pass
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 4, 40)

    now = timezone.localtime().strftime('%Y%m-%d_%H%M')
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="students_report_{now}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_students_pdf(request):
    """Export filtered students list as a PDF report table."""
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_RIGHT
    from offers.views import ARABIC_FONT, _prepare_arabic

    user = request.user
    qs = filter_by_branch(
        Student.objects.select_related('contact', 'branch'), user, 'branch', perm='view_student'
    )
    qs = _apply_student_filters(qs, request.GET)
    qs = qs.order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=1*cm, leftMargin=1*cm,
                            topMargin=1*cm, bottomMargin=1*cm)
    elements = []

    font_name = ARABIC_FONT
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleAr', parent=styles['Normal'], fontName=font_name, fontSize=16, alignment=TA_RIGHT, spaceAfter=12, leading=22)
    header_style = ParagraphStyle('HeaderAr', parent=styles['Normal'], fontName=font_name, fontSize=9, alignment=TA_RIGHT,
                                  textColor=colors.white, leading=14)
    cell_style = ParagraphStyle('CellAr', parent=styles['Normal'], fontName=font_name, fontSize=8, alignment=TA_RIGHT, leading=12)

    elements.append(Paragraph(_prepare_arabic('تقرير الطلاب'), title_style))
    elements.append(Spacer(1, 0.3*cm))

    headers = ['الاسم', 'الجوال', 'المستوى', 'طريقة التواصل', 'الفرع', 'تاريخ الإنشاء']
    table_data = [[Paragraph(_prepare_arabic(h), header_style) for h in headers]]

    for student in qs:
        row = [
            Paragraph(_prepare_arabic(student.get_full_name()), cell_style),
            Paragraph(_prepare_arabic(student.get_mobile() or '-'), cell_style),
            Paragraph(_prepare_arabic(student.level), cell_style),
            Paragraph(_prepare_arabic(student.get_preferred_contact_display()), cell_style),
            Paragraph(_prepare_arabic(str(student.branch)), cell_style),
            Paragraph(_prepare_arabic(student.created_at.strftime('%Y-%m-%d %H:%M') if student.created_at else ''), cell_style),
        ]
        table_data.append(row)

    page_w = landscape(A4)[0] - 2*cm
    col_widths = [page_w*0.22, page_w*0.18, page_w*0.14, page_w*0.16, page_w*0.18, page_w*0.12]

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563EB')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, -1), font_name),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F1F5F9')]),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor('#2563EB')),
        ('INNERGRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#CBD5E1')),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 0.5*cm))
    count_style = ParagraphStyle('CountAr', parent=styles['Normal'], fontName=font_name, fontSize=10, alignment=TA_RIGHT)
    elements.append(Paragraph(_prepare_arabic(f'إجمالي النتائج: {qs.count()}'), count_style))

    doc.build(elements)

    now = timezone.localtime().strftime('%Y%m-%d_%H%M')
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="students_report_{now}.pdf"'
    return response
